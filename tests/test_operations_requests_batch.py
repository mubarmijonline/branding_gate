"""
What David's Operations team asked for, after testing with it.

1. A client's approval is where Operations' work starts, and nobody was told:
   the item just appeared on Approved Items for whoever opened the page.
2. Approved Items: every request looked the same, and the table was a flat
   list of items. Each request now has its own colour, and the table is one
   row per request that expands to its items.
3. The exports carried columns nobody on site needs -- an internal id, the
   rent/sell type, what the item cost us, the supplier's email, the approval
   date -- and the costs should not reach a supplier at all. Which columns go
   out is now chosen, and those six are off unless asked for.
5. My Expenses: the supplier dropdown was empty for Operations. The list came
   from an endpoint only Finance could read, so it answered 403 and the page
   said nothing. It now opens to anyone who records an expense, defaults to
   Transportation, and "Other" saves the name typed rather than the word.

Rollback-based, like the other route tests here.
"""

import io
import os
import unittest

import MySQLdb
import MySQLdb.cursors

import branding_gate
import fixtures

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


class _RollbackConnection:
    def __init__(self, raw): self.raw = raw
    def commit(self): pass
    def close(self): pass


class _Harness(unittest.TestCase):

    def setUp(self):
        self.raw = MySQLdb.connect(host="localhost", user="ps", passwd="Aa@123456",
                                   db="branding_gate", charset="utf8mb4", use_unicode=True)
        self.raw.autocommit(False)
        self.original = branding_gate.connection
        self._real_notify = branding_gate.notify_users
        self.addCleanup(self._undo)
        branding_gate.connection = lambda: (_RollbackConnection(self.raw), self._cursor())
        branding_gate.app.config['TESTING'] = True
        self.sent = []
        branding_gate.notify_users = lambda ids, title, body, link=None: (
            self.sent.append({'to': sorted(int(i) for i in (ids or [])), 'title': title,
                              'body': body, 'link': link}) or len(ids or []))

    def _undo(self):
        branding_gate.notify_users = self._real_notify
        branding_gate.connection = self.original
        self.raw.rollback()
        self.raw.close()

    def _cursor(self):
        return self.raw.cursor(MySQLdb.cursors.DictCursor)

    def _client_for(self, user_id):
        perms, role_code = branding_gate.load_permissions(user_id)
        client = branding_gate.app.test_client()
        with client.session_transaction() as flask_session:
            flask_session.update({'user_id': user_id, 'mobile': 'm', 'email': 'e',
                                  'username': 'u', 'name': 'Tester', 'roles': [role_code],
                                  'perms': perms, 'role_code': role_code})
        return client

    def _user_with_role(self, role_code):
        cur = self._cursor()
        cur.execute("""SELECT u.id FROM user u JOIN rbac_role r ON r.id = u.rbac_role_id
                       WHERE r.code = %s ORDER BY u.id LIMIT 1""", (role_code,))
        row = cur.fetchone()
        cur.close()
        if not row:
            self.skipTest('no %s account' % role_code)
        return row['id']

    def _request_with_item(self, approval_status='pending'):
        cur = self._cursor()
        client_id = fixtures.ensure_client(cur)
        cur.execute("""INSERT INTO sales_request (client_id, title, start_date, end_date, created_by,
                                                  items_count, owner_user_id)
                       VALUES (%s, 'Ops batch probe', '2026-10-10', '2026-10-12', 'probe', 1, 1)""",
                    (client_id,))
        request_id = cur.lastrowid
        cur.execute("""INSERT INTO sales_request_items
                           (request_id, name, qty, unit, cost_per_item, sell_per_item,
                            total_cost, total_sell, approval_status, sell_type, rental_days)
                       VALUES (%s, 'Ops probe item', 10, 'pcs', 80, 120, 800, 1200, %s, 'rent', 1)""",
                    (request_id, approval_status))
        item_id = cur.lastrowid
        cur.close()
        return request_id, item_id


class ClientApprovalTellsOperationsTest(_Harness):

    def test_the_whole_team_hears_with_a_link_to_the_request(self):
        request_id, item_id = self._request_with_item('pending')
        response = self._client_for(1).post('/api/client-approval/items/%d/approve' % item_id,
                                            json={'approved': True})
        self.assertEqual(response.status_code, 200, response.data[:300])

        told = [n for n in self.sent if n['title'].startswith('Client approved')]
        self.assertEqual(len(told), 1, 'one notice per approval')
        cur = self._cursor()
        everyone = branding_gate.operations_team_ids(cur)
        cur.close()
        self.assertTrue(everyone, 'there is nobody in Operations to tell')
        self.assertEqual(told[0]['to'], sorted(everyone))
        self.assertEqual(told[0]['link'], '/approved-items?request=%d' % request_id)
        self.assertIn('Ops probe item', told[0]['body'])


class ExportColumnsTest(_Harness):

    HIDDEN = {'Item ID', 'Type', 'Cost/Unit', 'Total Cost', 'Supplier Email', 'Approval Date'}

    def _header(self, path):
        import openpyxl
        response = self._client_for(1).get(path)
        self.assertEqual(response.status_code, 200, response.data[:200])
        book = openpyxl.load_workbook(io.BytesIO(response.data))
        sheet = book[book.sheetnames[0]]
        return [cell.value for cell in next(sheet.iter_rows(max_row=1))]

    def test_the_six_are_off_by_default_in_every_export(self):
        request_id, _ = self._request_with_item('approved')
        for path in ('/api/operations/approved-items/export/by-request',
                     '/api/operations/approved-items/export/by-supplier',
                     '/api/operations/approved-items/export/request/%d' % request_id):
            with self.subTest(export=path):
                header = self._header(path)
                self.assertFalse(self.HIDDEN & set(header), header)
                self.assertIn('Item Name', header)
                self.assertIn('Supplier', header)

    def test_the_chosen_columns_are_the_columns_in_their_own_order(self):
        request_id, _ = self._request_with_item('approved')
        header = self._header('/api/operations/approved-items/export/request/%d'
                              '?cols=Total Cost,Item Name,Quantity' % request_id)
        self.assertEqual(header, ['Item Name', 'Quantity', 'Total Cost'])

    def test_unknown_or_empty_choices_fall_back_to_the_defaults(self):
        request_id, _ = self._request_with_item('approved')
        base = '/api/operations/approved-items/export/request/%d' % request_id
        default = self._header(base)
        self.assertEqual(self._header(base + '?cols=Nonsense'), default)
        self.assertEqual(self._header(base + '?cols='), default)

    def _rows(self, path):
        import openpyxl
        response = self._client_for(1).get(path)
        self.assertEqual(response.status_code, 200, response.data[:200])
        book = openpyxl.load_workbook(io.BytesIO(response.data))
        # Every sheet: the supplier export files an item with no supplier on
        # its own "Unassigned" sheet, last.
        out = []
        for name in book.sheetnames:
            rows = list(book[name].iter_rows(values_only=True))
            out.extend(dict(zip(rows[0], r)) for r in rows[1:])
        return out

    def test_the_setup_date_is_in_both_exports_by_default(self):
        # It lives in the request's template JSON, not in a column, which is
        # why it never reached the sheets.
        request_id, _ = self._request_with_item('approved')
        cur = self._cursor()
        cur.execute("""INSERT INTO sales_request_template_instances
                           (request_id, template_id, request_type, instance_order, template_data)
                       VALUES (%s, 1, 'Booth', 0, %s)""",
                    (request_id, '{"setup_date": "2026-10-09", "event_date": "2026-10-10"}'))
        cur.close()
        for path in ('/api/operations/approved-items/export/request/%d' % request_id,
                     '/api/operations/approved-items/export/by-supplier'):
            with self.subTest(export=path):
                mine = [r for r in self._rows(path) if r.get('Item Name') == 'Ops probe item']
                self.assertTrue(mine, 'the probe item is missing from the sheet')
                self.assertEqual(mine[0]['Setup Date'], '2026-10-09')

    def test_a_request_without_one_leaves_it_blank_not_null(self):
        request_id, _ = self._request_with_item('approved')
        cur = self._cursor()
        cur.execute("""INSERT INTO sales_request_template_instances
                           (request_id, template_id, request_type, instance_order, template_data)
                       VALUES (%s, 1, 'Event', 0, %s)""",
                    (request_id, '{"setup_date": null, "event_date": "2026-10-10"}'))
        cur.close()
        rows = self._rows('/api/operations/approved-items/export/request/%d' % request_id)
        self.assertIn(rows[0]['Setup Date'], (None, ''))

    def test_the_chooser_is_told_what_exists_and_what_is_off(self):
        body = self._client_for(1).get('/api/operations/approved-items/export/columns').get_json()
        self.assertEqual(set(body['hidden_by_default']), self.HIDDEN)
        self.assertTrue(self.HIDDEN <= set(body['columns']))



class SupplierReportSetupDateTest(_Harness):
    """
    The Supplier Report Generator exports from its own data, not the
    approved-items exports, so fixing those left it without the setup date.
    It reads the same single definition now (setup_date_sql).
    """

    def _assigned_probe(self):
        request_id, item_id = self._request_with_item('approved')
        cur = self._cursor()
        supplier_id = fixtures.ensure_suppliers(cur, 1)[0]
        cur.execute("UPDATE sales_request_items SET supplier_id = %s WHERE id = %s",
                    (supplier_id, item_id))
        cur.execute("""INSERT INTO sales_request_template_instances
                           (request_id, template_id, request_type, instance_order, template_data)
                       VALUES (%s, 1, 'Booth', 0, %s)""",
                    (request_id, '{"setup_date": "2026-10-09"}'))
        cur.close()
        return request_id

    def test_each_report_row_carries_its_setup_date(self):
        request_id = self._assigned_probe()
        items = self._client_for(1).get('/api/supplier-report').get_json().get('items') or []
        mine = [i for i in items if i.get('request_id') == request_id]
        self.assertTrue(mine, 'the probe item is not in the report')
        self.assertEqual(mine[0]['setup_date'], '2026-10-09')

    def test_the_report_export_button_writes_it(self):
        with open(os.path.join(ROOT, 'templates', 'approved_items.html'), encoding='utf-8') as handle:
            page = handle.read()
        start = page.index("$('#exportSupplierReportBtn').click")
        block = page[start:start + 2500]
        self.assertIn("'Setup Date'", block)
        self.assertIn("(item.setup_date || '')", block)

    def test_the_server_side_report_export_writes_it(self):
        import openpyxl
        self._assigned_probe()
        response = self._client_for(1).get('/api/supplier-report/export-excel')
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.data))
        header = [c.value for c in next(book[book.sheetnames[0]].iter_rows(max_row=1))]
        self.assertIn('Setup Date', header)

    def test_there_is_one_definition(self):
        with open(os.path.join(ROOT, 'branding_gate.py'), encoding='utf-8') as handle:
            source = handle.read()
        # the lookup itself appears once; everything else calls it
        self.assertEqual(source.count("FROM sales_request_template_instances ti"), 1)
        self.assertGreaterEqual(source.count("setup_date_sql("), 4)



class ExportFiltersTest(_Harness):
    """
    Exporting a chosen slice: certain requests, certain suppliers, a client,
    an event window, rent or sell. The filters ride on the export's query
    string and are applied once, in the fetch all three exports share.
    """

    def setUp(self):
        super().setUp()
        cur = self._cursor()
        client_id = fixtures.ensure_client(cur)
        cur.execute("""INSERT INTO supplier (supplier_name, status, added_by)
                       VALUES ('Filter Probe Supplier', 'Active', 'probe')""")
        self.supplier_id = cur.lastrowid
        self.requests = {}
        for key, start, sell_type, with_supplier in (('a', '2026-11-05', 'rent', True),
                                                     ('b', '2026-12-20', 'sell', False)):
            cur.execute("""INSERT INTO sales_request (client_id, title, start_date, end_date,
                                                      created_by, items_count, owner_user_id)
                           VALUES (%s, %s, %s, %s, 'probe', 1, 1)""",
                        (client_id, 'Filter probe ' + key, start, start))
            request_id = cur.lastrowid
            cur.execute("""INSERT INTO sales_request_items
                               (request_id, name, qty, unit, cost_per_item, sell_per_item,
                                total_cost, total_sell, approval_status, sell_type, rental_days,
                                supplier_id)
                           VALUES (%s, %s, 1, 'pcs', 10, 20, 10, 20, 'approved', %s, 1, %s)""",
                        (request_id, 'Filter item ' + key, sell_type,
                         self.supplier_id if with_supplier else None))
            self.requests[key] = request_id
        cur.execute("SELECT client_name FROM client WHERE id = %s", (client_id,))
        self.client_name = cur.fetchone()['client_name']
        cur.close()

    def _names(self, query):
        import openpyxl
        response = self._client_for(1).get('/api/operations/approved-items/export/by-request?' + query)
        if response.status_code == 404:
            return None
        self.assertEqual(response.status_code, 200, response.data[:200])
        book = openpyxl.load_workbook(io.BytesIO(response.data))
        names = set()
        for sheet in book.sheetnames:
            rows = list(book[sheet].iter_rows(values_only=True))
            at = rows[0].index('Item Name')
            names.update(r[at] for r in rows[1:])
        return {n for n in names if n and n.startswith('Filter item')}

    def _only_probes(self):
        return 'request_ids=%d,%d' % (self.requests['a'], self.requests['b'])

    def test_certain_requests(self):
        self.assertEqual(self._names('request_ids=%d' % self.requests['a']), {'Filter item a'})
        self.assertEqual(self._names(self._only_probes()), {'Filter item a', 'Filter item b'})

    def test_certain_suppliers_and_the_unassigned(self):
        self.assertEqual(self._names('supplier_ids=%d' % self.supplier_id), {'Filter item a'})
        self.assertEqual(self._names(self._only_probes() + '&supplier_ids=unassigned'), {'Filter item b'})
        self.assertEqual(self._names(self._only_probes() + '&supplier_ids=%d,unassigned' % self.supplier_id),
                         {'Filter item a', 'Filter item b'})

    def test_event_window_and_rent_or_sell(self):
        q = self._only_probes()
        self.assertEqual(self._names(q + '&start_from=2026-12-01'), {'Filter item b'})
        self.assertEqual(self._names(q + '&start_to=2026-11-30'), {'Filter item a'})
        self.assertEqual(self._names(q + '&sell_type=sell'), {'Filter item b'})

    def test_client_by_name(self):
        import urllib.parse
        q = self._only_probes() + '&clients=' + urllib.parse.quote(self.client_name)
        self.assertEqual(self._names(q), {'Filter item a', 'Filter item b'})
        self.assertIsNone(self._names(self._only_probes() + '&clients=Nobody%20Called%20This'))

    def test_filters_combine(self):
        q = 'supplier_ids=%d&sell_type=rent&start_from=2026-11-01&start_to=2026-11-30' % self.supplier_id
        self.assertEqual(self._names(q), {'Filter item a'})

    def test_an_unreadable_value_is_not_a_filter(self):
        self.assertEqual(self._names(self._only_probes() + '&start_from=not-a-date&sell_type=lease'),
                         {'Filter item a', 'Filter item b'})

    def test_nothing_matching_says_so(self):
        response = self._client_for(1).get('/api/operations/approved-items/export/by-request'
                                            '?request_ids=%d&sell_type=sell' % self.requests['a'])
        self.assertEqual(response.status_code, 404)
        self.assertIn('match this export', response.get_json()['error'])

    def test_the_supplier_export_takes_the_same_filters(self):
        import openpyxl
        response = self._client_for(1).get('/api/operations/approved-items/export/by-supplier'
                                            '?supplier_ids=%d' % self.supplier_id)
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.data))
        self.assertEqual(book.sheetnames, ['Filter Probe Supplier'])

    def test_the_page_offers_and_previews_them(self):
        with open(os.path.join(ROOT, 'templates', 'approved_items.html'), encoding='utf-8') as handle:
            page = handle.read()
        self.assertIn('Export options', page)
        for picker in ('aiPickRequests', 'aiPickSuppliers', 'aiPickClients', 'aiPickFrom',
                       'aiPickTo', 'aiPickSellType', 'aiMatchCount'):
            self.assertIn('id="%s"' % picker, page)
        for param in ('request_ids=', 'supplier_ids=', 'clients=', 'start_from=', 'start_to=', 'sell_type='):
            self.assertIn(param, page)
        # Filters last for the visit only; columns are remembered.
        self.assertNotIn('localStorage.setItem(AI_FILTERS', page)
        self.assertIn("$('#aiColumnsSave').prop('disabled', n === 0);", page)


class ExpenseSuppliersTest(_Harness):

    def test_operations_can_read_the_supplier_list(self):
        # Azer's role. This was a 403, and the dropdown said nothing.
        response = self._client_for(self._user_with_role('operations_member')).get('/api/finance/suppliers')
        self.assertEqual(response.status_code, 200, response.data[:200])
        self.assertTrue(response.get_json().get('success'))


class PagesTest(unittest.TestCase):

    def _page(self, name):
        with open(os.path.join(ROOT, 'templates', name), encoding='utf-8') as handle:
            return handle.read()

    def test_the_expense_supplier_defaults_to_transportation_and_other_is_typed(self):
        page = self._page('my_expenses.html')
        self.assertIn('<option value="__transport" selected>Transportation</option>', page)
        self.assertIn("supplierName = 'Transportation';", page)
        # "Other" saves what was typed, and cannot be saved blank.
        self.assertIn("supplierName = ($('#expenseSupplierName').val() || '').trim();", page)
        self.assertIn('for "Other"', page)
        self.assertNotIn('<option value="">Select Supplier</option>', page)

    def test_each_request_has_its_own_colour_in_both_views(self):
        page = self._page('approved_items.html')
        self.assertIn('function requestColor(id)', page)
        self.assertIn("'<div class=\"ai-req\" style=\"' + requestStyle(g.request_id) + '\">'", page)
        self.assertIn("row.setAttribute('style', requestStyle(g.request_id));", page)

    def test_the_table_is_one_row_per_request_that_expands(self):
        page = self._page('approved_items.html')
        self.assertIn('var groups = groupByRequest(allItems);', page)
        self.assertIn("row.child('<div class=\"ai-rq-items\"", page)
        self.assertIn('g.items.map(aiLineHtml)', page)
        self.assertNotIn('rowGroup:', page)

    def test_every_export_carries_the_chosen_columns(self):
        page = self._page('approved_items.html')
        self.assertIn('id="aiColumnsBtn"', page)
        self.assertEqual(page.count('ai-export-link'), 3)       # two links and the handler
        self.assertIn("exportUrl('/api/operations/approved-items/export/request/' + rid)", page)

    def test_the_notice_opens_its_own_request(self):
        page = self._page('approved_items.html')
        self.assertIn("new URLSearchParams(window.location.search).get('request')", page)



class OperationsMenuTest(_Harness):
    """Approved Items is in the Operations menu, for whoever may open it."""

    def _operations_menu(self, user_id):
        html = self._client_for(user_id).get('/home').get_data(as_text=True)
        start = html.find('id="operationsDropdown"')
        if start == -1:
            return None
        end = html.find('</li>', start)
        return html[start:end]

    def test_david_finds_it_in_the_operations_menu(self):
        menu = self._operations_menu(self._user_with_role('operations_manager'))
        self.assertIsNotNone(menu, 'no Operations menu at all')
        self.assertIn('href="/approved-items"', menu)
        # after Costing, before Workflow Timeline: the step that follows costing
        self.assertLess(menu.index('/approved-items'), menu.index('Workflow Timeline'))

    def test_it_is_drawn_only_for_those_who_can_open_it(self):
        with open(os.path.join(ROOT, 'templates', 'main.html'), encoding='utf-8') as handle:
            page = handle.read()
        at = page.index("{{ url_for('approved_items_page') }}")
        self.assertIn("{% if 'approved_item.view' in _p %}", page[at - 200:at])


if __name__ == '__main__':
    unittest.main()
